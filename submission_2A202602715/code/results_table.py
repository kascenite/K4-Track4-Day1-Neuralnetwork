"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền experiments.xlsx từ
templates/experiment_table_template.xlsx (giữ nguyên 4 sheet, tên cột và các cột công thức).
"""
from __future__ import annotations

import json
import math
from pathlib import Path

COLUMNS = ["exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs",
           "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss",
           "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
           "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes"]
FORMULA_COLUMNS = ("step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise")
MAX_ROWS = 60   # công thức của sheet Seeds/Summary chỉ tham chiếu dòng 2..61

LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    with open(path, "w") as f:
        json.dump({k: result[k] for k in ("cfg", "history", "summary")}, f, indent=1)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi *.json trong results_dir (sắp theo exp_id)."""
    out = []
    for p in sorted(Path(results_dir).glob("*.json")):
        with open(p) as f:
            out.append(json.load(f))
    return out


def _num(x, nd=None):
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return None
    return round(x, nd) if nd is not None else x


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một kết quả → một dòng bảng (khoá trùng tên cột). Chỉ truyền eval_scores cho baseline và cấu hình cuối."""
    cfg, s = result["cfg"], result["summary"]
    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg["description"],
        loss=LOSS_NAMES[cfg["loss"]], optimizer=OPT_NAMES[cfg["optimizer"]], lr=cfg["lr"],
        weight_decay=cfg["weight_decay"], batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(str(h) for h in cfg["hidden"]), dropout=cfg["dropout"],
        clip_norm="none" if cfg["clip_norm"] is None else round(cfg["clip_norm"], 4),
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        step0_loss=_num(s["step0_loss"], 4), best_val_loss=_num(s["best_val_loss"], 4),
        best_epoch=s["best_epoch"], final_train_loss=_num(s["final_train_loss"], 4),
        final_val_loss=_num(s["final_val_loss"], 4), val_acc=_num(s["val_acc"], 4),
        val_macro_f1=_num(s["val_macro_f1"], 4), time_per_epoch_s=_num(s["time_per_epoch_s"], 3),
        peak_mem_MB=_num(s["peak_mem_MB"], 1), diverged="Y" if s["diverged"] else "N",
        eval_acc=None, eval_macro_f1=None, figure_file=f"figures/{cfg['exp_id']}.png", notes=notes,
    )
    if eval_scores is not None:
        row["eval_acc"] = round(eval_scores["accuracy"], 4)
        row["eval_macro_f1"] = round(eval_scores["macro_f1"], 4)
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền rows vào sheet Experiments (từ dòng 2), exp_id baseline vào sheet Seeds (cột A),
    nhận xét vào sheet Summary (cột H, theo group). Không ghi đè các cột công thức."""
    import openpyxl

    assert len(rows) <= MAX_ROWS, f"tối đa {MAX_ROWS} dòng (công thức mẫu chỉ phủ dòng 2..61)"
    assert len({r["exp_id"] for r in rows}) == len(rows), "exp_id phải duy nhất"
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    header = {c.value: c.column for c in ws[1] if c.value is not None}
    assert all(c in header for c in COLUMNS), "mẫu thiếu cột"
    for i in range(2, MAX_ROWS + 2):      # xoá dòng ví dụ của mẫu (chỉ ô nhập, giữ công thức)
        for c in COLUMNS:
            ws.cell(row=i, column=header[c]).value = None
    for i, row in enumerate(rows, start=2):
        for c in COLUMNS:
            ws.cell(row=i, column=header[c]).value = row.get(c)

    if seed_ids is not None:
        ss = wb["Seeds"]
        assert len(seed_ids) <= 5
        for i in range(2, 7):
            ss.cell(row=i, column=1).value = seed_ids[i - 2] if i - 2 < len(seed_ids) else None

    if summary_notes:
        sm = wb["Summary"]
        for r in range(2, sm.max_row + 1):
            g = sm.cell(row=r, column=1).value
            if g in summary_notes:
                sm.cell(row=r, column=8).value = summary_notes[g]
    wb.save(out_path)
